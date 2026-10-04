"""
Text extraction from uploaded resume files.

Deliberately free of Django and of the database: it takes a file object and
returns a string, so it can be tested directly against the sample files in
sample_data/.

Only PDF and DOCX are handled. A scanned PDF that contains images rather than
a text layer will extract to nothing -- that is treated as a failure with a
message the candidate can act on, not as a silent empty resume.
"""

import os

import pdfplumber
from docx import Document


class ResumeParseError(Exception):
    """Raised when a file cannot be turned into usable text."""


# A resume shorter than this almost certainly means extraction failed -- a
# scanned image, or a PDF whose text layer is missing.
MIN_USABLE_CHARACTERS = 50


def extract_text(file_object, filename):
    """
    Return the text of a PDF or DOCX file.

    `file_object` is anything with a file-like interface, which is what Django
    hands over for an upload. Raises ResumeParseError with a readable message
    on any failure, so the caller never has to inspect a library-specific
    exception.
    """
    extension = os.path.splitext(filename)[1].lower()

    if extension == ".pdf":
        text = _extract_pdf(file_object)
    elif extension == ".docx":
        text = _extract_docx(file_object)
    else:
        raise ResumeParseError(f"صيغة غير مدعومة: {extension or 'غير معروفة'}")

    text = text.strip()
    if len(text) < MIN_USABLE_CHARACTERS:
        raise ResumeParseError(
            "تعذّر استخراج نص كافٍ من الملف. "
            "إذا كان الملف صورة ممسوحة ضوئيًا، ارفع نسخة نصية منه."
        )

    return text


def _extract_pdf(file_object):
    try:
        lines = []
        with pdfplumber.open(file_object) as pdf:
            for page in pdf.pages:
                # extract_text() returns None for a page with no text layer.
                lines.append(page.extract_text() or "")
        return "\n".join(lines)
    except ResumeParseError:
        raise
    except Exception as exc:
        # pdfplumber raises a variety of types on a damaged file. The caller
        # only needs to know that this one could not be read.
        raise ResumeParseError(f"تعذّرت قراءة ملف PDF: {exc}") from exc


def _extract_docx(file_object):
    try:
        document = Document(file_object)

        parts = [paragraph.text for paragraph in document.paragraphs]

        # Many resumes lay out skills and dates in tables, and python-docx does
        # not include table text in `paragraphs`. Skipping them would drop
        # exactly the content the matching engine most needs.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)

        return "\n".join(parts)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"تعذّرت قراءة ملف DOCX: {exc}") from exc
