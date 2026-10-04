"""
Phase 4 tests: text extraction, upload validation, and the resume manager.

The parsing tests run against the real files in sample_data/, so they exercise
the same code path a candidate's upload takes rather than a mock of it.
"""

import json
import os
import shutil
import tempfile

from django.core import mail
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from accounts.models import CompanyProfile, SeekerProfile, User

from .models import Resume
from .parsing import ResumeParseError, extract_text

PASSWORD = "pw-for-tests-9931"
SAMPLE_DIR = os.path.join(settings.BASE_DIR, "sample_data")

# Uploads are written to a temporary directory that is removed afterwards, so
# a test run never leaves files behind in MEDIA_ROOT.
TEMP_MEDIA = tempfile.mkdtemp(prefix="ats-test-media-")


def sample_bytes(name):
    with open(os.path.join(SAMPLE_DIR, name), "rb") as handle:
        return handle.read()


def sample_upload(name, upload_name=None):
    return SimpleUploadedFile(upload_name or name, sample_bytes(name))


def make_seeker(email="seeker@example.com", name="محمد عثمان"):
    user = User.objects.create_user(email=email, password=PASSWORD, role=User.Role.JOB_SEEKER)
    return SeekerProfile.objects.create(user=user, full_name=name)


def make_company(email="co@example.com", name="سوداتل"):
    user = User.objects.create_user(email=email, password=PASSWORD, role=User.Role.COMPANY)
    return CompanyProfile.objects.create(user=user, name=name)


class TextExtractionTests(TestCase):
    """Runs against the committed sample files, not mocks."""

    def test_pdf_text_is_extracted(self):
        text = extract_text(sample_upload("backend_developer_cv.pdf"),
                            "backend_developer_cv.pdf")
        self.assertIn("Mohammed Osman", text)
        self.assertIn("Django and Python", text)
        self.assertIn("PostgreSQL", text)

    def test_docx_text_is_extracted(self):
        text = extract_text(sample_upload("frontend_developer_cv.docx"),
                            "frontend_developer_cv.docx")
        self.assertIn("Salma Idris", text)
        self.assertIn("Next.js and Redux", text)

    def test_the_same_content_extracts_from_either_format(self):
        pdf = extract_text(sample_upload("backend_developer_cv.pdf"), "a.pdf")
        docx = extract_text(sample_upload("backend_developer_cv.docx"), "a.docx")
        for marker in ["Mohammed Osman", "Django and Python", "University of Khartoum"]:
            with self.subTest(marker=marker):
                self.assertIn(marker, pdf)
                self.assertIn(marker, docx)

    def test_docx_table_cells_are_included(self):
        """Resumes often lay skills out in tables, which paragraphs alone miss."""
        from docx import Document
        import io

        document = Document()
        document.add_paragraph("Fatima Ali, Data Analyst working out of Khartoum")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Wrote complex SQL queries against a warehouse"
        table.rows[0].cells[1].text = "Built dashboards for operations teams"

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)

        text = extract_text(buffer, "tabled.docx")
        self.assertIn("complex SQL queries", text)
        self.assertIn("Built dashboards", text)

    def test_multi_page_pdf_extracts_every_page(self):
        """
        _extract_pdf loops over pdf.pages; this is the fixture that proves the
        loop is needed. resume.pdf runs to two pages, and content from the
        second one has to survive.
        """
        import pdfplumber

        with pdfplumber.open(os.path.join(SAMPLE_DIR, "resume.pdf")) as pdf:
            self.assertGreater(len(pdf.pages), 1)

        text = extract_text(sample_upload("resume.pdf"), "resume.pdf")
        self.assertIn("Nour Abdelrahman", text)          # first page
        self.assertIn("Available on request", text)      # last page

    def test_the_long_resume_produces_many_chunks(self):
        from matching.chunking import split_into_chunks

        text = extract_text(sample_upload("resume.pdf"), "resume.pdf")
        chunks = split_into_chunks(text)
        # A real CV should yield far more evidence lines than the short fixtures.
        self.assertGreater(len(chunks), 20)
        self.assertTrue(all(len(c.split()) >= 4 for c in chunks))

    def test_unsupported_extension_is_rejected(self):
        with self.assertRaises(ResumeParseError):
            extract_text(SimpleUploadedFile("cv.txt", b"x" * 200), "cv.txt")

    def test_corrupt_pdf_raises_a_readable_error_not_a_library_exception(self):
        broken = SimpleUploadedFile("broken.pdf", b"%PDF-1.4 this is not a real pdf")
        with self.assertRaises(ResumeParseError):
            extract_text(broken, "broken.pdf")

    def test_pdf_with_no_text_layer_is_treated_as_a_failure(self):
        """A scanned resume extracts to nothing, which must not pass silently."""
        # Blank out the text-showing operators: the PDF still parses, but no
        # words come out, which is what a scanned page looks like.
        blank = sample_bytes("backend_developer_cv.pdf").replace(b"Tj", b"  ")
        empty_pdf = SimpleUploadedFile("scan.pdf", blank)
        with self.assertRaises(ResumeParseError):
            extract_text(empty_pdf, "scan.pdf")


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class UploadValidationTests(TestCase):
    def setUp(self):
        self.seeker = make_seeker()
        self.client.force_login(self.seeker.user)
        self.url = reverse("applications:resume_upload")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def test_unsupported_extension_is_rejected_before_saving(self):
        response = self.client.post(
            self.url, {"file": SimpleUploadedFile("cv.txt", b"plain text resume")}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("file", response.context["form"].errors)
        self.assertFalse(Resume.objects.exists())

    def test_oversized_file_is_rejected(self):
        oversized = SimpleUploadedFile(
            "huge.pdf", b"x" * (settings.RESUME_MAX_UPLOAD_BYTES + 1)
        )
        response = self.client.post(self.url, {"file": oversized})
        self.assertEqual(response.status_code, 200)
        self.assertIn("file", response.context["form"].errors)
        self.assertFalse(Resume.objects.exists())

    def test_empty_file_is_rejected(self):
        response = self.client.post(self.url, {"file": SimpleUploadedFile("cv.pdf", b"")})
        self.assertEqual(response.status_code, 200)
        self.assertIn("file", response.context["form"].errors)
        self.assertFalse(Resume.objects.exists())

    def test_a_file_just_under_the_cap_is_accepted_by_the_form(self):
        from .forms import ResumeUploadForm

        ok = SimpleUploadedFile("cv.pdf", b"x" * (settings.RESUME_MAX_UPLOAD_BYTES - 1))
        form = ResumeUploadForm(files={"file": ok})
        self.assertTrue(form.is_valid(), form.errors)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class ResumeProcessingTests(TestCase):
    """End to end: upload, parse, embed, store."""

    def setUp(self):
        self.seeker = make_seeker()
        self.client.force_login(self.seeker.user)
        self.url = reverse("applications:resume_upload")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def test_pdf_upload_is_parsed_and_embedded(self):
        response = self.client.post(
            self.url, {"file": sample_upload("backend_developer_cv.pdf")}
        )
        self.assertRedirects(response, reverse("applications:resume_list"))

        resume = Resume.objects.get()
        self.assertEqual(resume.parse_status, Resume.ParseStatus.DONE)
        self.assertEqual(resume.parse_error, "")
        self.assertIn("Django and Python", resume.parsed_text)
        self.assertTrue(resume.doc_embedding)
        self.assertTrue(resume.chunk_embeddings)

    def test_docx_upload_is_parsed_and_embedded(self):
        self.client.post(self.url, {"file": sample_upload("frontend_developer_cv.docx")})
        resume = Resume.objects.get()
        self.assertEqual(resume.parse_status, Resume.ParseStatus.DONE)
        self.assertIn("Next.js and Redux", resume.parsed_text)

    def test_stored_embeddings_load_back_as_usable_vectors(self):
        from matching import embeddings

        self.client.post(self.url, {"file": sample_upload("backend_developer_cv.pdf")})
        resume = Resume.objects.get()

        doc_vector = embeddings.vector_from_json(resume.doc_embedding)
        chunk_texts, chunk_matrix = embeddings.chunks_from_json(resume.chunk_embeddings)

        self.assertEqual(doc_vector.shape, (384,))
        self.assertGreater(len(chunk_texts), 0)
        self.assertEqual(chunk_matrix.shape[1], 384)
        self.assertEqual(len(chunk_texts), chunk_matrix.shape[0])
        # Every stored chunk must be a real line from the resume.
        for chunk in chunk_texts:
            self.assertIn(chunk.split()[0], resume.parsed_text)

    def test_a_broken_pdf_is_recorded_as_failed_and_does_not_500(self):
        response = self.client.post(
            self.url, {"file": SimpleUploadedFile("broken.pdf", b"%PDF-1.4 garbage")}
        )
        self.assertRedirects(response, reverse("applications:resume_list"))

        resume = Resume.objects.get()
        self.assertEqual(resume.parse_status, Resume.ParseStatus.FAILED)
        self.assertTrue(resume.parse_error)
        self.assertFalse(resume.is_active)

    def test_a_failed_resume_is_never_made_active(self):
        self.client.post(self.url, {"file": SimpleUploadedFile("broken.pdf", b"%PDF garbage")})
        self.assertIsNone(
            Resume.objects.filter(is_active=True).first()
        )

    def test_first_successful_upload_becomes_active(self):
        self.client.post(self.url, {"file": sample_upload("backend_developer_cv.pdf")})
        self.assertTrue(Resume.objects.get().is_active)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class ActiveResumeTests(TestCase):
    """Exactly one resume per seeker is active, and activating is scoped."""

    def setUp(self):
        self.seeker = make_seeker()
        self.other = make_seeker("other@example.com", "سلمى إدريس")
        self.client.force_login(self.seeker.user)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def _resume(self, seeker=None, status=Resume.ParseStatus.DONE, active=False):
        return Resume.objects.create(
            seeker=seeker or self.seeker,
            file=SimpleUploadedFile("cv.pdf", b"stub"),
            original_filename="cv.pdf",
            parse_status=status,
            is_active=active,
        )

    def test_activating_one_deactivates_the_previous(self):
        first = self._resume(active=True)
        second = self._resume()

        second.activate()

        first.refresh_from_db()
        second.refresh_from_db()
        self.assertFalse(first.is_active)
        self.assertTrue(second.is_active)

    def test_activating_does_not_touch_another_seekers_resumes(self):
        theirs = self._resume(seeker=self.other, active=True)
        mine = self._resume()

        mine.activate()

        theirs.refresh_from_db()
        self.assertTrue(theirs.is_active)

    def test_activate_view_requires_post(self):
        resume = self._resume()
        url = reverse("applications:resume_activate", args=[resume.pk])
        self.assertEqual(self.client.get(url).status_code, 405)

    def test_cannot_activate_a_failed_resume(self):
        resume = self._resume(status=Resume.ParseStatus.FAILED)
        self.client.post(reverse("applications:resume_activate", args=[resume.pk]))
        resume.refresh_from_db()
        self.assertFalse(resume.is_active)

    def test_cannot_activate_another_seekers_resume(self):
        theirs = self._resume(seeker=self.other)
        response = self.client.post(
            reverse("applications:resume_activate", args=[theirs.pk])
        )
        self.assertEqual(response.status_code, 404)
        theirs.refresh_from_db()
        self.assertFalse(theirs.is_active)

    def test_active_resume_helper_ignores_unparsed_rows(self):
        from . import services

        self._resume(status=Resume.ParseStatus.FAILED, active=True)
        self.assertIsNone(services.active_resume_for(self.seeker))


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class ResumeAccessTests(TestCase):
    def setUp(self):
        self.seeker = make_seeker()
        self.other = make_seeker("other@example.com", "سلمى إدريس")
        self.company = make_company()

        self.resume = Resume.objects.create(
            seeker=self.seeker,
            file=SimpleUploadedFile("cv.pdf", b"stub"),
            original_filename="cv.pdf",
            parse_status=Resume.ParseStatus.DONE,
        )

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def test_company_cannot_reach_the_resume_manager(self):
        self.client.force_login(self.company.user)
        for name in ("applications:resume_list", "applications:resume_upload"):
            with self.subTest(view=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 403)

    def test_anonymous_visitor_is_redirected_to_login(self):
        response = self.client.get(reverse("applications:resume_list"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("accounts:login"), response.url)

    def test_seeker_sees_only_their_own_resumes(self):
        Resume.objects.create(
            seeker=self.other,
            file=SimpleUploadedFile("theirs.pdf", b"stub"),
            original_filename="theirs.pdf",
        )
        self.client.force_login(self.seeker.user)
        response = self.client.get(reverse("applications:resume_list"))
        self.assertEqual(list(response.context["resumes"]), [self.resume])

    def test_cannot_delete_another_seekers_resume(self):
        self.client.force_login(self.other.user)
        response = self.client.post(
            reverse("applications:resume_delete", args=[self.resume.pk])
        )
        self.assertEqual(response.status_code, 404)
        self.assertTrue(Resume.objects.filter(pk=self.resume.pk).exists())

    def test_delete_requires_post(self):
        self.client.force_login(self.seeker.user)
        url = reverse("applications:resume_delete", args=[self.resume.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertTrue(Resume.objects.filter(pk=self.resume.pk).exists())

        self.client.post(url)
        self.assertFalse(Resume.objects.filter(pk=self.resume.pk).exists())


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class EndToEndScoringTests(TestCase):
    """
    The whole pipeline in one test: an uploaded file becomes text, becomes
    vectors, and scores against a real Job row.

    Phase 5 wires this into the apply flow. This test proves the two halves fit
    together before any of that exists.
    """

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from jobs.models import Job
        from matching import services as matching_services

        self.company = make_company()
        self.job = Job.objects.create(
            company=self.company,
            title="Backend Developer",
            description=(
                "Build and maintain the server-side services behind our payments "
                "product. You will design HTTP APIs and model data in a relational "
                "database."
            ),
            required_skills=(
                "building REST APIs with Django and Python\n"
                "optimising slow PostgreSQL queries\n"
                "automated testing and continuous integration"
            ),
            location="الخرطوم",
            status=Job.Status.OPEN,
        )
        built = matching_services.build_job_embeddings(
            self.job.description, self.job.skill_lines()
        )
        self.job.description_embedding = built["description_embedding"]
        self.job.skills_embedding = built["skills_embedding"]
        self.job.save()

    def _upload(self, filename):
        seeker = make_seeker(f"{filename}@example.com", filename)
        self.client.force_login(seeker.user)
        self.client.post(
            reverse("applications:resume_upload"), {"file": sample_upload(filename)}
        )
        return Resume.objects.get(seeker=seeker)

    def test_matching_resume_outscores_an_unrelated_one(self):
        from matching import services as matching_services

        backend = self._upload("backend_developer_cv.pdf")
        chef = self._upload("chef_cv.pdf")

        good = matching_services.score_resume_against_job(backend, self.job)
        bad = matching_services.score_resume_against_job(chef, self.job)

        self.assertGreater(good["score"], bad["score"])
        self.assertGreater(good["coverage"], bad["coverage"])

    def test_evidence_lines_come_from_the_uploaded_file(self):
        from matching import services as matching_services

        backend = self._upload("backend_developer_cv.pdf")
        result = matching_services.score_resume_against_job(backend, self.job)

        for entry in result["skill_results"]:
            with self.subTest(skill=entry["skill"]):
                self.assertIn(entry["evidence"], backend.parsed_text)


# --------------------------------------------------------------------------
# Phase 5: applying, the pipeline, stage changes and notes
# --------------------------------------------------------------------------


def make_job(company, **overrides):
    from jobs.models import Job

    data = {
        "company": company,
        "title": "Backend Developer",
        "description": "Build and maintain server-side services for payments.",
        "required_skills": (
            "building REST APIs with Django and Python\n"
            "optimising slow PostgreSQL queries"
        ),
        "location": "الخرطوم",
        "status": Job.Status.OPEN,
    }
    data.update(overrides)
    return Job.objects.create(**data)


def stub_resume(seeker, status=None):
    """
    A Resume row without real embeddings.

    Used by the tests that exercise stage changes and access rules, where the
    score is irrelevant. Keeping the model out of those tests is what stops the
    suite from loading a transformer to check a 403.
    """
    return Resume.objects.create(
        seeker=seeker,
        file=SimpleUploadedFile("cv.pdf", b"stub"),
        original_filename="cv.pdf",
        parse_status=status or Resume.ParseStatus.DONE,
        is_active=True,
    )


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class ApplyFlowTests(TestCase):
    """The real path: a parsed resume scored against an embedded job."""

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from jobs.services import embed_job
        from .models import Application

        self.company = make_company()
        self.job = make_job(self.company)
        embed_job(self.job)

        self.seeker = make_seeker()
        self.client.force_login(self.seeker.user)
        self.client.post(
            reverse("applications:resume_upload"),
            {"file": sample_upload("backend_developer_cv.pdf")},
        )
        self.Application = Application

    def test_applying_creates_a_scored_application(self):
        response = self.client.post(
            reverse("jobs:apply", args=[self.job.pk]), {"cover_letter": "Keen to join."}
        )
        self.assertRedirects(response, reverse("applications:my_applications"))

        application = self.Application.objects.get()
        self.assertEqual(application.seeker, self.seeker)
        self.assertEqual(application.stage, self.Application.Stage.APPLIED)
        self.assertEqual(application.cover_letter, "Keen to join.")

        # A relevant CV against a matching job must score above zero on all three.
        self.assertGreater(application.match_score, 0)
        self.assertGreater(application.semantic_score, 0)
        self.assertGreater(application.coverage_score, 0)

    def test_skill_results_are_stored_with_evidence(self):
        self.client.post(reverse("jobs:apply", args=[self.job.pk]))
        application = self.Application.objects.get()

        breakdown = application.skill_breakdown
        self.assertEqual(len(breakdown), 2)
        for entry in breakdown:
            with self.subTest(skill=entry["skill"]):
                self.assertIn("similarity", entry)
                self.assertIn("matched", entry)
                # Every evidence line must come from the candidate's own resume.
                self.assertIn(entry["evidence"], application.resume.parsed_text)

    def test_applying_opens_the_stage_timeline(self):
        self.client.post(reverse("jobs:apply", args=[self.job.pk]))
        application = self.Application.objects.get()

        history = list(application.history.all())
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0].from_stage, "")
        self.assertEqual(history[0].to_stage, self.Application.Stage.APPLIED)

    def test_applying_twice_is_refused(self):
        self.client.post(reverse("jobs:apply", args=[self.job.pk]))
        response = self.client.post(reverse("jobs:apply", args=[self.job.pk]), follow=True)

        self.assertEqual(self.Application.objects.count(), 1)
        self.assertEqual(response.status_code, 200)

    def test_the_unique_constraint_is_enforced_at_the_database(self):
        """The friendly message is a nicety; the constraint is the guarantee."""
        from django.db import IntegrityError

        self.client.post(reverse("jobs:apply", args=[self.job.pk]))
        existing = self.Application.objects.get()

        with self.assertRaises(IntegrityError):
            self.Application.objects.create(
                job=self.job, seeker=self.seeker, resume=existing.resume
            )

    def test_a_job_created_outside_the_form_is_embedded_on_first_apply(self):
        """Seeded and admin-created jobs must not silently score everyone zero."""
        unembedded = make_job(self.company, title="Seeded Role")
        self.assertEqual(unembedded.skills_embedding, "")

        self.client.post(reverse("jobs:apply", args=[unembedded.pk]))

        unembedded.refresh_from_db()
        self.assertTrue(unembedded.skills_embedding)
        self.assertGreater(self.Application.objects.get(job=unembedded).match_score, 0)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class ApplyGuardTests(TestCase):
    """Who may apply, and to what. No scoring involved, so no model is loaded."""

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.company = make_company()
        self.seeker = make_seeker()
        self.job = make_job(self.company)

    def test_a_seeker_with_no_resume_cannot_apply(self):
        from .models import Application
        from . import services

        with self.assertRaises(services.ApplicationError):
            services.apply_to_job(self.seeker, self.job)
        self.assertFalse(Application.objects.exists())

    def test_a_seeker_whose_resume_failed_parsing_cannot_apply(self):
        from . import services

        stub_resume(self.seeker, status=Resume.ParseStatus.FAILED)
        with self.assertRaises(services.ApplicationError):
            services.apply_to_job(self.seeker, self.job)

    def test_a_draft_job_does_not_accept_applications(self):
        from jobs.models import Job
        from . import services

        stub_resume(self.seeker)
        draft = make_job(self.company, status=Job.Status.DRAFT)
        with self.assertRaises(services.ApplicationError):
            services.apply_to_job(self.seeker, draft)

    def test_a_closed_job_does_not_accept_applications(self):
        from jobs.models import Job
        from . import services

        stub_resume(self.seeker)
        closed = make_job(self.company, status=Job.Status.CLOSED)
        with self.assertRaises(services.ApplicationError):
            services.apply_to_job(self.seeker, closed)

    def test_a_company_cannot_apply_to_a_job(self):
        self.client.force_login(self.company.user)
        response = self.client.get(reverse("jobs:apply", args=[self.job.pk]))
        self.assertEqual(response.status_code, 403)

    def test_an_anonymous_visitor_is_sent_to_login(self):
        response = self.client.get(reverse("jobs:apply", args=[self.job.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("accounts:login"), response.url)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class PipelineAccessTests(TestCase):
    """Company A must never see company B's candidates."""

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application

        self.owner = make_company("owner@example.com", "سوداتل")
        self.intruder = make_company("intruder@example.com", "زين")
        self.seeker = make_seeker()

        self.job = make_job(self.owner)
        self.application = Application.objects.create(
            job=self.job, seeker=self.seeker, resume=stub_resume(self.seeker)
        )
        self.Application = Application

    def test_another_company_cannot_open_the_pipeline(self):
        self.client.force_login(self.intruder.user)
        response = self.client.get(reverse("jobs:pipeline", args=[self.job.pk]))
        self.assertEqual(response.status_code, 404)

    def test_another_company_cannot_open_a_candidate(self):
        self.client.force_login(self.intruder.user)
        response = self.client.get(
            reverse("applications:detail", args=[self.application.pk])
        )
        self.assertEqual(response.status_code, 404)

    def test_a_seeker_cannot_open_the_pipeline(self):
        self.client.force_login(self.seeker.user)
        self.assertEqual(
            self.client.get(reverse("jobs:pipeline", args=[self.job.pk])).status_code, 403
        )

    def test_a_seeker_cannot_open_the_candidate_review_screen(self):
        """Even their own application: this screen is the company's view."""
        self.client.force_login(self.seeker.user)
        self.assertEqual(
            self.client.get(
                reverse("applications:detail", args=[self.application.pk])
            ).status_code,
            403,
        )

    def test_the_owning_company_can_open_both(self):
        self.client.force_login(self.owner.user)
        self.assertEqual(
            self.client.get(reverse("jobs:pipeline", args=[self.job.pk])).status_code, 200
        )
        self.assertEqual(
            self.client.get(
                reverse("applications:detail", args=[self.application.pk])
            ).status_code,
            200,
        )

    def test_a_seeker_sees_only_their_own_applications(self):
        other = make_seeker("other@example.com", "سلمى")
        self.Application.objects.create(
            job=self.job, seeker=other, resume=stub_resume(other)
        )

        self.client.force_login(self.seeker.user)
        response = self.client.get(reverse("applications:my_applications"))
        self.assertEqual(list(response.context["applications"]), [self.application])


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class RecruiterNoteTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application

        self.company = make_company()
        self.seeker = make_seeker()
        self.application = Application.objects.create(
            job=make_job(self.company), seeker=self.seeker, resume=stub_resume(self.seeker)
        )
        self.url = reverse("applications:add_note", args=[self.application.pk])
        self.client.force_login(self.company.user)

    def test_a_note_is_saved_against_its_author(self):
        self.client.post(self.url, {"body": "خبرة قوية في Django", "rating": 4})

        note = self.application.notes.get()
        self.assertEqual(note.body, "خبرة قوية في Django")
        self.assertEqual(note.rating, 4)
        self.assertEqual(note.author, self.company.user)

    def test_the_rating_is_optional(self):
        self.client.post(self.url, {"body": "ملاحظة بدون تقييم"})
        self.assertIsNone(self.application.notes.get().rating)

    def test_an_empty_note_is_refused(self):
        self.client.post(self.url, {"body": "   "})
        self.assertFalse(self.application.notes.exists())

    def test_a_rating_outside_one_to_five_is_refused(self):
        self.client.post(self.url, {"body": "ملاحظة", "rating": 9})
        self.assertFalse(self.application.notes.exists())

    def test_another_company_cannot_add_a_note(self):
        intruder = make_company("intruder@example.com", "زين")
        self.client.force_login(intruder.user)
        self.assertEqual(self.client.post(self.url, {"body": "x"}).status_code, 404)

    def test_adding_a_note_requires_post(self):
        self.assertEqual(self.client.get(self.url).status_code, 405)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class PipelineBoardTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application

        self.company = make_company()
        self.job = make_job(self.company)
        self.Application = Application
        self.client.force_login(self.company.user)

    def _application(self, name, score, stage="applied"):
        seeker = make_seeker(f"{name}@example.com", name)
        return self.Application.objects.create(
            job=self.job,
            seeker=seeker,
            resume=stub_resume(seeker),
            match_score=score,
            stage=stage,
            skill_results=json.dumps(
                [
                    {"skill": "a", "similarity": 0.9, "matched": True, "evidence": "x"},
                    {"skill": "b", "similarity": 0.1, "matched": False, "evidence": "y"},
                ]
            ),
        )

    def test_the_board_has_all_six_stages_in_order(self):
        response = self.client.get(reverse("jobs:pipeline", args=[self.job.pk]))
        stages = [column["stage"] for column in response.context["columns"]]
        self.assertEqual(
            stages, ["applied", "screening", "interview", "offer", "hired", "rejected"]
        )

    def test_candidates_are_grouped_into_their_stage(self):
        self._application("alpha", 0.9, stage="interview")
        self._application("beta", 0.8, stage="applied")

        response = self.client.get(reverse("jobs:pipeline", args=[self.job.pk]))
        columns = {c["stage"]: c["applications"] for c in response.context["columns"]}

        self.assertEqual([a.seeker.full_name for a in columns["interview"]], ["alpha"])
        self.assertEqual([a.seeker.full_name for a in columns["applied"]], ["beta"])
        self.assertEqual(columns["hired"], [])

    def test_candidates_are_sorted_by_score_descending(self):
        self._application("low", 0.20)
        self._application("high", 0.90)
        self._application("mid", 0.55)

        response = self.client.get(reverse("jobs:pipeline", args=[self.job.pk]))
        applied = next(
            c["applications"] for c in response.context["columns"] if c["stage"] == "applied"
        )
        self.assertEqual([a.seeker.full_name for a in applied], ["high", "mid", "low"])

    def test_the_card_counts_matched_skills(self):
        application = self._application("alpha", 0.7)
        self.assertEqual(application.matched_skill_count, 1)
        self.assertEqual(application.total_skill_count, 2)

    def test_malformed_skill_results_do_not_break_the_board(self):
        """Bad JSON in one row must not take the whole pipeline down."""
        seeker = make_seeker("broken@example.com", "broken")
        self.Application.objects.create(
            job=self.job,
            seeker=seeker,
            resume=stub_resume(seeker),
            skill_results="{not valid json",
        )
        response = self.client.get(reverse("jobs:pipeline", args=[self.job.pk]))
        self.assertEqual(response.status_code, 200)

    def test_the_board_does_not_issue_one_query_per_candidate(self):
        """
        The view the spec warns N+1s badly.

        Each card prints the candidate's name, which lives on a related row.
        Without select_related the page cost grows with the number of
        applicants; this pins it flat.
        """
        from . import services

        for i in range(3):
            self._application(f"cand{i}", 0.5)
        with self.assertNumQueries(1):
            services.pipeline_columns(self.job)

        for i in range(3, 12):
            self._application(f"cand{i}", 0.5)
        with self.assertNumQueries(1):
            services.pipeline_columns(self.job)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class CandidateDetailTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application

        self.company = make_company()
        self.seeker = make_seeker()
        resume = stub_resume(self.seeker)
        resume.parsed_text = "Built REST APIs with Django and Python at a telecom"
        resume.save()

        self.application = Application.objects.create(
            job=make_job(self.company),
            seeker=self.seeker,
            resume=resume,
            match_score=0.71,
            semantic_score=0.62,
            coverage_score=0.80,
            skill_results=json.dumps(
                [
                    {
                        "skill": "React",
                        "similarity": 0.71,
                        "matched": True,
                        "evidence": "built single-page apps with Next.js and Redux",
                    },
                    {
                        "skill": "Kubernetes",
                        "similarity": 0.19,
                        "matched": False,
                        "evidence": "deployed to Heroku",
                    },
                ]
            ),
        )
        self.client.force_login(self.company.user)

    def test_the_evidence_table_renders_every_skill_with_its_line(self):
        response = self.client.get(
            reverse("applications:detail", args=[self.application.pk])
        )
        self.assertEqual(response.status_code, 200)

        content = response.content.decode()
        for fragment in [
            "React",
            "0.710",
            "built single-page apps with Next.js and Redux",
            "Kubernetes",
            "0.190",
            "deployed to Heroku",
        ]:
            with self.subTest(fragment=fragment):
                self.assertIn(fragment, content)

    def test_unmatched_skills_still_show_their_closest_line(self):
        """'No, and here is the nearest thing we found' is the useful answer."""
        response = self.client.get(
            reverse("applications:detail", args=[self.application.pk])
        )
        self.assertIn("أقرب سطر وُجد", response.content.decode())

    def test_the_three_scores_are_shown(self):
        response = self.client.get(
            reverse("applications:detail", args=[self.application.pk])
        )
        content = response.content.decode()
        self.assertIn("0.71", content)
        self.assertIn("0.62", content)
        self.assertIn("0.80", content)

    def test_the_parsed_resume_text_is_shown(self):
        response = self.client.get(
            reverse("applications:detail", args=[self.application.pk])
        )
        self.assertIn("Built REST APIs with Django and Python", response.content.decode())


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class StageProgressionTests(TestCase):
    """
    The recruiter sends an intent, never a stage name.

    These tests pin that contract: accepting moves one step along the forward
    path, rejecting is reachable from anywhere, and neither the sequence nor
    the destination can be influenced by the request.
    """

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application

        self.company = make_company()
        self.seeker = make_seeker()
        self.job = make_job(self.company)
        self.application = Application.objects.create(
            job=self.job, seeker=self.seeker, resume=stub_resume(self.seeker)
        )
        self.Application = Application
        self.url = reverse("applications:stage_moved", args=[self.application.pk])
        self.client.force_login(self.company.user)

    def _post(self, action, ajax=True, **extra):
        headers = {"HTTP_X_REQUESTED_WITH": "XMLHttpRequest"} if ajax else {}
        return self.client.post(self.url, {"action": action, **extra}, **headers)

    # -- the sequence itself ------------------------------------------------

    def test_next_stage_for_walks_the_forward_path(self):
        from . import services

        self.assertEqual(services.next_stage_for("applied"), "screening")
        self.assertEqual(services.next_stage_for("screening"), "interview")
        self.assertEqual(services.next_stage_for("interview"), "offer")
        self.assertEqual(services.next_stage_for("offer"), "hired")

    def test_terminal_stages_have_no_next(self):
        from . import services

        self.assertIsNone(services.next_stage_for("hired"))
        self.assertIsNone(services.next_stage_for("rejected"))

    def test_accepting_repeatedly_walks_the_whole_pipeline(self):
        expected = ["screening", "interview", "offer", "hired"]
        for stage in expected:
            with self.subTest(stage=stage):
                response = self._post("accept")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["stage"], stage)

        self.application.refresh_from_db()
        self.assertEqual(self.application.stage, "hired")

    def test_accepting_a_hired_candidate_is_refused(self):
        for _ in range(4):
            self._post("accept")

        response = self._post("accept")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["ok"])

    # -- rejection ----------------------------------------------------------

    def test_rejection_works_from_any_stage(self):
        self._post("accept")   # screening
        self._post("accept")   # interview

        response = self._post("reject")
        self.assertEqual(response.json()["stage"], "rejected")
        self.assertFalse(response.json()["can_advance"])
        self.assertTrue(response.json()["is_rejected"])

    def test_a_rejected_candidate_cannot_be_advanced(self):
        self._post("reject")
        response = self._post("accept")
        self.assertEqual(response.status_code, 400)

    def test_rejecting_twice_is_refused(self):
        self._post("reject")
        self.assertEqual(self._post("reject").status_code, 400)

    # -- history ------------------------------------------------------------

    def test_every_move_records_a_history_row(self):
        self._post("accept")
        self._post("accept")
        self._post("reject")

        moves = [(h.from_stage, h.to_stage) for h in self.application.history.all()]
        self.assertEqual(
            moves,
            [("applied", "screening"), ("screening", "interview"), ("interview", "rejected")],
        )

    def test_the_note_is_stored_on_the_history_row(self):
        self._post("accept", note="مرشح قوي")
        self.assertEqual(self.application.history.latest("created_at").note, "مرشح قوي")

    def test_the_move_is_attributed_to_the_logged_in_recruiter(self):
        self._post("accept")
        self.assertEqual(
            self.application.history.latest("created_at").changed_by, self.company.user
        )

    # -- the contract -------------------------------------------------------

    def test_a_stage_name_in_the_request_is_ignored(self):
        """
        The whole point of the redesign: the client cannot choose a stage.

        Posting `stage=hired` alongside an accept must still move the
        candidate exactly one step, to screening.
        """
        response = self._post("accept", stage="hired")
        self.assertEqual(response.json()["stage"], "screening")

    def test_an_unknown_action_is_refused(self):
        response = self._post("promote_to_ceo")
        self.assertEqual(response.status_code, 400)
        self.application.refresh_from_db()
        self.assertEqual(self.application.stage, "applied")

    def test_the_response_carries_what_the_card_needs_to_redraw(self):
        payload = self._post("accept").json()
        for key in [
            "ok", "stage", "stage_label", "can_advance",
            "next_stage_label", "is_rejected", "changed_at",
        ]:
            with self.subTest(key=key):
                self.assertIn(key, payload)
        self.assertEqual(payload["next_stage_label"], "المقابلة")

    # -- access -------------------------------------------------------------

    def test_a_seeker_cannot_move_a_stage(self):
        self.client.force_login(self.seeker.user)
        self.assertEqual(self._post("accept").status_code, 403)
        self.application.refresh_from_db()
        self.assertEqual(self.application.stage, "applied")

    def test_another_company_cannot_move_a_stage(self):
        self.client.force_login(make_company("intruder@example.com", "زين").user)
        self.assertEqual(self._post("accept").status_code, 404)

    def test_the_endpoint_requires_post(self):
        self.assertEqual(self.client.get(self.url).status_code, 405)

    def test_it_still_works_without_javascript(self):
        """A plain form POST gets a redirect instead of JSON."""
        response = self._post("accept", ajax=False)
        self.assertEqual(response.status_code, 302)
        self.application.refresh_from_db()
        self.assertEqual(self.application.stage, "screening")


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class NotificationTriggerTests(TestCase):
    """
    The four things the system tells people about.

    Tested here rather than in notifications/ because what matters is that the
    actions fire them, not that notify() works -- that has its own tests.
    """

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        from .models import Application
        from notifications.models import Notification

        self.company = make_company()
        self.seeker = make_seeker()
        self.job = make_job(self.company)
        self.application = Application.objects.create(
            job=self.job, seeker=self.seeker, resume=stub_resume(self.seeker)
        )
        self.Application = Application
        self.Notification = Notification
        mail.outbox = []

    def _seeker_notifications(self):
        return self.Notification.objects.filter(user=self.seeker.user)

    def _company_notifications(self):
        return self.Notification.objects.filter(user=self.company.user)

    # -- stage changes ------------------------------------------------------

    def test_a_stage_change_notifies_the_candidate(self):
        from . import services

        services.change_stage(self.application, "screening", self.company.user)

        notification = self._seeker_notifications().get()
        self.assertEqual(notification.title, "تم تحديث حالة طلبك")
        self.assertIn("الفرز المبدئي", notification.body)

    def test_a_stage_change_emails_the_candidate(self):
        from . import services

        services.change_stage(self.application, "interview", self.company.user)

        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [self.seeker.user.email])
        self.assertIn("المقابلة", mail.outbox[0].subject)

    def test_rejection_gets_its_own_wording(self):
        from . import services

        services.reject_application(self.application, self.company.user)

        notification = self._seeker_notifications().get()
        self.assertEqual(notification.title, "لم يتم قبول طلبك")
        self.assertIn("لم يُقبل", notification.body)
        self.assertEqual(len(mail.outbox), 1)

    def test_a_recruiter_note_travels_with_the_notification(self):
        from . import services

        services.change_stage(
            self.application, "offer", self.company.user, note="نرجو تأكيد الموعد"
        )
        self.assertIn("نرجو تأكيد الموعد", self._seeker_notifications().get().body)

    def test_a_no_op_stage_change_notifies_nobody(self):
        """Resubmitting the form must not send a second email."""
        from . import services

        services.change_stage(self.application, "applied", self.company.user)

        self.assertFalse(self._seeker_notifications().exists())
        self.assertEqual(len(mail.outbox), 0)

    def test_the_company_is_not_notified_about_its_own_stage_change(self):
        from . import services

        services.change_stage(self.application, "screening", self.company.user)
        self.assertFalse(self._company_notifications().exists())

    # -- new applications ---------------------------------------------------

    def test_a_new_application_notifies_the_company(self):
        from jobs.services import embed_job
        from . import services

        embed_job(self.job)
        other = make_seeker("newcomer@example.com", "سلمى إدريس")
        self.client.force_login(other.user)
        self.client.post(
            reverse("applications:resume_upload"),
            {"file": sample_upload("backend_developer_cv.pdf")},
        )
        services.apply_to_job(other, self.job)

        notification = self._company_notifications().get()
        self.assertEqual(notification.title, "طلب توظيف جديد")
        self.assertIn("سلمى إدريس", notification.body)
        self.assertIn("/applications/", notification.link)

    # -- resume parsing -----------------------------------------------------

    def test_a_failed_resume_notifies_the_candidate(self):
        self.client.force_login(self.seeker.user)
        self.client.post(
            reverse("applications:resume_upload"),
            {"file": SimpleUploadedFile("broken.pdf", b"%PDF-1.4 garbage")},
        )
        notification = self._seeker_notifications().get()
        self.assertEqual(notification.title, "تعذّرت معالجة سيرتك الذاتية")
        self.assertIn("broken.pdf", notification.body)

    def test_a_successful_resume_notifies_nobody(self):
        self.client.force_login(self.seeker.user)
        self.client.post(
            reverse("applications:resume_upload"),
            {"file": sample_upload("backend_developer_cv.pdf")},
        )
        self.assertFalse(self._seeker_notifications().exists())

    # -- resilience ---------------------------------------------------------

    def test_a_broken_notifier_does_not_roll_back_the_stage_change(self):
        """
        The action is what matters; the notification is a side effect.

        If the notification table were unavailable, a recruiter's decision
        must still stick.
        """
        from unittest.mock import patch
        from . import services

        with patch(
            "notifications.services.Notification.objects.create",
            side_effect=Exception("db is down"),
        ):
            services.change_stage(self.application, "screening", self.company.user)

        self.application.refresh_from_db()
        self.assertEqual(self.application.stage, "screening")
        self.assertEqual(self.application.history.count(), 1)
